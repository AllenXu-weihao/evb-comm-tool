"""测试报告导出 — 每条用例导出仿真行 + 实测行。"""

from __future__ import annotations

import csv
import math
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from analysis.sign_compare import (
    ber_channel_pairs_for_mode,
    ber_compare_label_for_mode,
    ber_meas_out_indices_for_mode,
    bit_label,
    chip_value_to_bit,
    compare_channels_for_mode,
    sim_value_to_bit,
)
from analysis.aggregate import (
    AttemptRecord,
    CaseRecord,
    analyze_by_snr,
    frame_was_tested,
    has_measured,
    apply_bit_compare,
    sim_row_display_infobits,
)
from csv_loader import ROW_KIND_MEASURED, ROW_KIND_SIM, detect_coded_bit_prefix
from protocol.codec import raw_to_row_voltage


def _bit_cell(row: dict[str, str], key: str) -> str:
    return str(row.get(key, "")).strip()


def measured_row_ber_mismatch_pairs(
    sim: dict[str, str],
    meas: dict[str, str],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> list[tuple[int, int]]:
    """返回 BER 误码通道：(实测 Out 下标, 仿真 CodeBit 下标)。"""
    prefix = detect_coded_bit_prefix(fieldnames)
    if not prefix:
        return []
    bad: list[tuple[int, int]] = []
    for out_i, cb_i in ber_channel_pairs_for_mode(int(matrix_mode)):
        cb_key = f"{prefix}[{cb_i}]"
        out_key = f"Out[{out_i}]"
        if cb_key not in fieldnames or out_key not in fieldnames:
            continue
        sim_val = _bit_cell(sim, cb_key)
        meas_out = str(meas.get(out_key, "")).strip()
        if not sim_val or not meas_out:
            continue
        try:
            meas_bit = chip_value_to_bit(float(meas_out))
            ref_bit = int(float(sim_val))
        except ValueError:
            continue
        if meas_bit != ref_bit:
            bad.append((out_i, cb_i))
    return bad


def measured_row_ber_mismatch_indices(
    sim: dict[str, str],
    meas: dict[str, str],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> list[int]:
    """返回实测行 BER 误码的 Out 下标列表（兼容旧接口）。"""
    return [out_i for out_i, _ in measured_row_ber_mismatch_pairs(
        sim, meas, fieldnames, matrix_mode=matrix_mode
    )]


def _snr_from_row(row: dict[str, str]) -> float | None:
    raw = str(row.get("EbNo_dB", "")).strip()
    if not raw:
        return None
    try:
        return float(raw)
    except ValueError:
        return None


@dataclass
class SnrExportCounts:
    """测后导出行按 SNR 计数（便于与折线图验算）。"""

    sim_frames: int = 0
    meas_rows: int = 0


def export_row_counts_by_snr(rows: list[dict[str, str]]) -> dict[float, SnrExportCounts]:
    """按 EbNo_dB 统计仿真帧数与实测行数（含多 Attempt）。"""
    by_snr: dict[float, SnrExportCounts] = {}
    i = 0
    n = len(rows)
    while i < n:
        row = rows[i]
        if str(row.get("RowKind", "")).strip() not in (ROW_KIND_SIM, "仿真行"):
            i += 1
            continue
        snr = _snr_from_row(row)
        if snr is None:
            i += 1
            continue
        c = by_snr.setdefault(snr, SnrExportCounts())
        c.sim_frames += 1
        j = i + 1
        while j < n and str(rows[j].get("RowKind", "")).strip() in (
            ROW_KIND_MEASURED,
            "实测行",
        ):
            c.meas_rows += 1
            j += 1
        by_snr[snr] = c
        i = j
    return by_snr


def sim_row_reference_error_bits(
    sim: dict[str, str],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> int:
    """仿真行参考误码：Sim Out 硬判 vs 指定 CodeBit（按模式配对下标）。"""
    prefix = detect_coded_bit_prefix(fieldnames)
    if not prefix:
        return 0
    err = 0
    for out_i, cb_i in ber_channel_pairs_for_mode(matrix_mode):
        out_key = f"Out[{out_i}]"
        cb_key = f"{prefix}[{cb_i}]"
        out_s = str(sim.get(out_key, "")).strip()
        cb_s = str(sim.get(cb_key, "")).strip()
        if not out_s or not cb_s:
            continue
        try:
            out_bit = chip_value_to_bit(float(out_s))
            cb_bit = int(float(cb_s))
        except ValueError:
            continue
        if out_bit != cb_bit:
            err += 1
    return err


def reference_ber_stats_by_snr_from_exports(
    rows: list[dict[str, str]],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> dict[float, tuple[int, int, int]]:
    """从 ok 导出行仿真行按 SNR 统计参考 BER：(error_bits, channel_bits, frame_count)。"""
    channels = len(ber_meas_out_indices_for_mode(matrix_mode))
    by_snr: dict[float, tuple[int, int, int]] = {}
    for row in rows:
        if str(row.get("RowKind", "")).strip() not in (ROW_KIND_SIM, "仿真行"):
            continue
        snr = _snr_from_row(row)
        if snr is None:
            continue
        err = sim_row_reference_error_bits(row, fieldnames, matrix_mode=matrix_mode)
        eb, cb, fc = by_snr.get(snr, (0, 0, 0))
        by_snr[snr] = (eb + err, cb + channels, fc + 1)
    return by_snr


def ber_formula_lines(matrix_mode: int) -> list[str]:
    """可视化页展示的 BER / 参考 BER 计算公式说明。"""
    from ui.chart_settings import get_frames_per_snr

    cap = get_frames_per_snr()
    ch = compare_channels_for_mode(matrix_mode)
    ch_desc = ber_compare_label_for_mode(matrix_mode)
    return [
        f"实测 BER = Σ(ber 表误码 bit) ÷ (图表帧数 × {ch})",
        f"  · 分子：{ch_desc} 不一致 bit 数（ber 表每帧末次误码 Attempt）",
        f"  · 分母：ok 表每 SNR 至多 {cap} 帧（FrameIdx 1..{cap}；1 仿真 + 1 实测/帧）",
        "参考 BER = Σ(Sim Out 硬判 bit ≠ 指定 CodeBit) ÷ (图表仿真行数 × "
        f"{ch})",
        f"  · 分母：与实测 BER 相同（图表聚合时按每 SNR 前 {cap} 帧裁剪）",
    ]


def build_snr_audit_rows(
    points: list,
    *,
    full_rows: list[dict[str, str]],
    ok_rows: list[dict[str, str]],
    ber_rows: list[dict[str, str]],
    fieldnames: list[str],
    matrix_mode: int,
) -> list[dict[str, Any]]:
    """折线图 vs 测后 ok/ber 表逐 SNR 验算行（供可视化分析表格）。"""
    from csv_loader import trim_ok_rows_for_chart

    chart_rows = trim_ok_rows_for_chart(ok_rows)
    meas_exp = measured_ber_stats_by_snr_from_exports(
        chart_rows, fieldnames, matrix_mode=matrix_mode
    )
    ref_exp = reference_ber_stats_by_snr_from_exports(
        chart_rows, fieldnames, matrix_mode=matrix_mode
    )
    full_c = export_row_counts_by_snr(full_rows)
    ok_c = export_row_counts_by_snr(ok_rows)
    chart_c = export_row_counts_by_snr(chart_rows)
    ber_c = export_row_counts_by_snr(ber_rows)

    point_by_snr = {float(p.snr_db): p for p in points}
    all_snrs = sorted(
        set(point_by_snr)
        | set(meas_exp)
        | set(ref_exp)
        | set(full_c)
    )

    rows_out: list[dict[str, Any]] = []
    for snr in all_snrs:
        p = point_by_snr.get(snr)
        eb_m, cb_m, chart_fc = meas_exp.get(snr, (0, 0, 0))
        export_meas_ber = (
            eb_m / cb_m if cb_m > 0 else float("nan")
        )
        chart_meas_ber = p.channel_ber if p and p.tested_count > 0 else float("nan")
        meas_ok = (
            p is not None
            and cb_m > 0
            and p.error_bits == eb_m
            and p.channel_bits == cb_m
            and p.tested_count == chart_fc
            and (
                math.isnan(chart_meas_ber)
                or math.isnan(export_meas_ber)
                or abs(chart_meas_ber - export_meas_ber) < 1e-12
            )
        )

        eb_r, cb_r, ref_fc = ref_exp.get(snr, (0, 0, 0))
        export_ref_ber = eb_r / cb_r if cb_r > 0 else float("nan")
        chart_ref_ber = (
            p.reference_ber
            if p and p.reference_ber is not None and p.reference_channel_bits > 0
            else float("nan")
        )
        ref_ok = (
            p is not None
            and cb_r > 0
            and p.reference_error_bits == eb_r
            and p.reference_channel_bits == cb_r
            and p.tested_count == chart_fc
            and (
                p.reference_ber is None
                or math.isnan(chart_ref_ber)
                or math.isnan(export_ref_ber)
                or abs(chart_ref_ber - export_ref_ber) < 1e-12
            )
        )

        fc = full_c.get(snr, SnrExportCounts())
        oc = ok_c.get(snr, SnrExportCounts())
        chart_cnt = chart_c.get(snr, SnrExportCounts())

        ber_cnt = ber_c.get(snr, SnrExportCounts())
        frame_audit_ok = (
            chart_cnt.sim_frames == chart_cnt.meas_rows == chart_fc
            and chart_fc <= oc.meas_rows
            and chart_fc <= oc.sim_frames
        )

        rows_out.append(
            {
                "snr_db": snr,
                "chart_frames": chart_fc,
                "ber_error_frames": ber_cnt.sim_frames,
                "meas_error_bits": eb_m,
                "meas_channel_bits": cb_m,
                "meas_ber_chart": chart_meas_ber,
                "meas_ber_export": export_meas_ber,
                "meas_match": meas_ok,
                "ref_error_bits": eb_r,
                "ref_channel_bits": cb_r,
                "ref_frames": ref_fc,
                "ref_ber_chart": chart_ref_ber,
                "ref_ber_export": export_ref_ber,
                "ref_match": ref_ok,
                "full_meas_rows": fc.meas_rows,
                "ok_meas_rows": oc.meas_rows,
                "ok_sim_rows": oc.sim_frames,
                "frame_audit_ok": frame_audit_ok,
                "ber_frames": ber_cnt.sim_frames,
            }
        )
    return rows_out


def _parse_iter_cell(row: dict[str, str]) -> int:
    raw = str(row.get("IterNum", "")).strip()
    if not raw:
        return 0
    try:
        return int(float(raw))
    except ValueError:
        return 0


def iteration_stats_by_snr_from_chart_exports(
    chart_rows: list[dict[str, str]],
) -> dict[float, tuple[int, int, int]]:
    """从 ok 图表裁剪子集按 SNR 统计迭代：(meas_sum, ref_sum, frame_count)。"""
    by_snr: dict[float, tuple[int, int, int]] = {}
    i = 0
    n = len(chart_rows)
    while i < n:
        row = chart_rows[i]
        if str(row.get("RowKind", "")).strip() not in (ROW_KIND_SIM, "仿真行"):
            i += 1
            continue
        sim = row
        snr = _snr_from_row(sim)
        if snr is None:
            i += 1
            continue
        j = i + 1
        if j >= n or str(chart_rows[j].get("RowKind", "")).strip() not in (
            ROW_KIND_MEASURED,
            "实测行",
        ):
            i += 1
            continue
        meas = chart_rows[j]
        ms, rs, fc = by_snr.get(snr, (0, 0, 0))
        by_snr[snr] = (
            ms + _parse_iter_cell(meas),
            rs + _parse_iter_cell(sim),
            fc + 1,
        )
        i = j + 1
    return by_snr


def iteration_formula_lines() -> list[str]:
    """可视化页展示的平均迭代次数验算公式。"""
    from ui.chart_settings import get_frames_per_snr

    cap = get_frames_per_snr()
    return [
        f"实测平均迭代次数 = Σ(ok 表实测行 IterNum) ÷ 图表帧数（每 SNR ≤ {cap}）",
        "  · 实测行：每帧末次收敛 Attempt（ok 表 1 仿真 + 1 实测/帧）",
        f"参考平均迭代次数 = Σ(ok 表仿真行 IterNum) ÷ 图表帧数（每 SNR ≤ {cap}）",
        "  · 仿真行：与 ok 表同帧配对的仿真 IterNum",
    ]


def build_iter_audit_rows(
    points: list,
    ok_rows: list[dict[str, str]],
) -> list[dict[str, Any]]:
    """折线图 vs ok 图表裁剪子集逐 SNR 平均迭代验算行。"""
    from csv_loader import trim_ok_rows_for_chart

    chart_rows = trim_ok_rows_for_chart(ok_rows)
    exp = iteration_stats_by_snr_from_chart_exports(chart_rows)
    point_by_snr = {float(p.snr_db): p for p in points}
    all_snrs = sorted(set(point_by_snr) | set(exp))

    rows_out: list[dict[str, Any]] = []
    for snr in all_snrs:
        p = point_by_snr.get(snr)
        meas_sum, ref_sum, fc = exp.get(snr, (0, 0, 0))
        export_meas_avg = meas_sum / fc if fc > 0 else float("nan")
        export_ref_avg = ref_sum / fc if fc > 0 else float("nan")
        chart_meas = (
            p.avg_iterations
            if p is not None and p.tested_count > 0
            else float("nan")
        )
        chart_ref = (
            p.avg_ref_iterations
            if p is not None
            and p.avg_ref_iterations is not None
            and p.tested_count > 0
            else float("nan")
        )
        meas_ok = (
            p is not None
            and fc > 0
            and p.tested_count == fc
            and math.isfinite(chart_meas)
            and math.isfinite(export_meas_avg)
            and abs(chart_meas - export_meas_avg) < 1e-6
        )
        ref_ok = (
            p is not None
            and fc > 0
            and p.tested_count == fc
            and math.isfinite(chart_ref)
            and math.isfinite(export_ref_avg)
            and abs(chart_ref - export_ref_avg) < 1e-6
        )
        rows_out.append(
            {
                "snr_db": snr,
                "chart_frames": fc,
                "meas_iter_sum": meas_sum,
                "ref_iter_sum": ref_sum,
                "meas_avg_chart": chart_meas,
                "meas_avg_export": export_meas_avg,
                "meas_match": meas_ok,
                "ref_avg_chart": chart_ref,
                "ref_avg_export": export_ref_avg,
                "ref_match": ref_ok,
            }
        )
    return rows_out


def measured_row_is_final_pass(
    sim: dict[str, str],
    meas: dict[str, str],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> bool:
    """通信收敛且相对仿真 CodeBit 无误码。"""
    from csv_loader import measured_row_is_nonconverged

    if measured_row_is_nonconverged(meas):
        return False
    return not measured_row_ber_mismatch_indices(
        sim, meas, fieldnames, matrix_mode=matrix_mode
    )


def measured_ber_stats_by_snr_from_exports(
    rows: list[dict[str, str]],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> dict[float, tuple[int, int, int]]:
    """从 ok 导出行按 SNR 统计实测 BER： (error_bits, channel_bits, frame_count)。

    与 analyze_by_snr 的图表口径一致；ber 表误码行应为本表中 error_bits>0 的子集。
    """
    from analysis.sign_compare import ber_meas_out_indices_for_mode

    channels = len(ber_meas_out_indices_for_mode(matrix_mode))
    by_snr: dict[float, tuple[int, int, int]] = {}
    i = 0
    n = len(rows)
    while i < n:
        row = rows[i]
        if str(row.get("RowKind", "")).strip() != ROW_KIND_SIM:
            i += 1
            continue
        sim = row
        snr_raw = str(sim.get("EbNo_dB", "")).strip()
        if not snr_raw:
            i += 1
            continue
        snr = float(snr_raw)
        j = i + 1
        if j >= n or str(rows[j].get("RowKind", "")).strip() != ROW_KIND_MEASURED:
            i += 1
            continue
        meas = rows[j]
        err_n = len(
            measured_row_ber_mismatch_indices(
                sim, meas, fieldnames, matrix_mode=matrix_mode
            )
        )
        eb, cb, fc = by_snr.get(snr, (0, 0, 0))
        by_snr[snr] = (eb + err_n, cb + channels, fc + 1)
        i = j + 1
    return by_snr


def filter_ber_error_rows(
    rows: list[dict[str, str]],
    fieldnames: list[str],
    *,
    matrix_mode: int,
) -> list[dict[str, str]]:
    """从 ok 表（1 仿真 + 1 实测/帧）只保留实测仍误码的帧。

    - 输入应为 ``build_export_rows(..., exclude_comm_errors=True)`` 的 ok 表
    - 收敛且无误码的帧整帧丢弃
    - 有误码的帧保留仿真行 + 该条实测行
    """
    from csv_loader import measured_row_is_nonconverged

    out: list[dict[str, str]] = []
    i = 0
    n = len(rows)
    while i < n:
        row = rows[i]
        kind = str(row.get("RowKind", "")).strip()
        if kind != ROW_KIND_SIM:
            i += 1
            continue
        sim = row
        j = i + 1
        if j >= n or str(rows[j].get("RowKind", "")).strip() != ROW_KIND_MEASURED:
            i += 1
            continue
        meas = rows[j]
        j += 1
        if measured_row_is_nonconverged(meas):
            i = j
            continue
        if measured_row_ber_mismatch_indices(
            sim, meas, fieldnames, matrix_mode=matrix_mode
        ):
            out.append(sim)
            out.append(meas)
        i = j
    return out


def _chip_bit_labels(
    currents: list[float],
    chip_bits: list[int] | None = None,
    *,
    mode: int = 1,
) -> list[str]:
    """CH1–4 判定 bit 标签（符号硬判）；优先 chip_bits。"""
    labels: list[str] = []
    for i in range(4):
        if chip_bits and i < len(chip_bits):
            labels.append(bit_label(int(chip_bits[i])))
        elif i < len(currents):
            labels.append(bit_label(chip_value_to_bit(currents[i])))
        else:
            labels.append("")
    return labels


def _fill_measured_infobits(
    row: dict,
    currents: list[float],
    chip_bits: list[int] | None = None,
    *,
    mode: int = 1,
) -> None:
    """实测行 InfoBit：按规则留空（不参与 BER）。"""
    del currents, chip_bits, mode
    for i in range(1, 5):
        row[f"InfoBit{i}"] = ""


def _blank_measured_fields() -> dict:
    row: dict = {
        "HardwareConverged": "",
        "FramePassed": "",
        "ChannelMatch": "",
        "Success": "",
        "Error": "",
        "RetryCount": "",
        "LowValueRetries": "",
        "Measured_Iterations": "",
        "SendHex": "",
        "RecvHex": "",
    }
    for i in range(1, 9):
        row[f"Chip_ICol{i}"] = ""
        row[f"RealAdj{i}"] = ""
        row[f"ChipSign{i}"] = ""
        row[f"SimSign{i}"] = ""
        row[f"InfoBit{i}"] = ""
        row[f"ChipBit{i}"] = ""
        row[f"Match_CH{i}"] = ""
        row[f"Chip_ADC{i}"] = ""
    return row


def _sim_row(tc: CaseRecord) -> dict:
    """仿真行：CSV In/Out、参考迭代；实测列留空。"""
    row: dict = {
        "RowKind": "仿真行",
        "Attempt": "",
        "CaseName": tc.name,
        "Mode": tc.mode,
        "SNR_dB": tc.snr_db,
        "FrameIdx": tc.frame_idx if tc.frame_idx is not None else "",
        "IterRef": tc.iter_ref if tc.iter_ref is not None else "",
    }
    row.update(_blank_measured_fields())

    for i in range(1, 13):
        ch = tc.channels[i - 1] if i - 1 < len(tc.channels) else ""
        row[f"In_CH{i}"] = ch
        row[f"In_Vrow{i}"] = f"{raw_to_row_voltage(int(ch)):.6f}" if ch != "" else ""

    sim = tc.sim_outputs or [0.0] * 8
    for i in range(1, 9):
        row[f"Sim_Out{i}"] = f"{sim[i - 1]:.6f}" if i - 1 < len(sim) else ""

    info = sim_row_display_infobits(tc) or []
    for i in range(1, 5):
        row[f"InfoBit{i}"] = bit_label(info[i - 1]) if i - 1 < len(info) else ""

    prefix = "CodedBit"
    for i in range(1, 9):
        sim_i = sim[i - 1] if i - 1 < len(sim) else 0.0
        row[f"{prefix}{i}"] = bit_label(sim_value_to_bit(sim_i))

    return row


def _measured_row_from_attempt(tc: CaseRecord, ar: AttemptRecord) -> dict:
    """单次尝试的实测导出行。"""
    row: dict = {
        "RowKind": "实测行",
        "Attempt": ar.attempt,
        "CaseName": tc.name,
        "Mode": tc.mode,
        "SNR_dB": tc.snr_db,
        "FrameIdx": tc.frame_idx if tc.frame_idx is not None else "",
        "IterRef": tc.iter_ref if tc.iter_ref is not None else "",
    }
    for i in range(1, 13):
        row[f"In_CH{i}"] = ""
        row[f"In_Vrow{i}"] = ""
    sim = tc.sim_outputs or [0.0] * 8
    for i in range(1, 9):
        row[f"Sim_Out{i}"] = f"{sim[i - 1]:.6f}" if i - 1 < len(sim) else ""
    row.update(
        {
            "HardwareConverged": "Y" if ar.hardware_converged else "N",
            "FramePassed": "Y" if ar.passed else ("N" if not ar.passed else ""),
            "ChannelMatch": f"{ar.match_count}/{compare_channels_for_mode(tc.mode)}",
            "Success": "Y",
            "Error": "" if ar.hardware_converged else "未收敛",
            "RetryCount": tc.retries,
            "LowValueRetries": tc.low_value_retries,
            "Measured_Iterations": ar.iteration_count if ar.iteration_count else "",
            "SendHex": tc.tx_hex,
            "RecvHex": ar.rx_hex,
        }
    )
    ch_results = {r["channel"]: r for r in ar.channel_results}
    for i in range(1, 9):
        cr = ch_results.get(i)
        if cr:
            row[f"Chip_ICol{i}"] = f"{cr['chip_current']:.6f}"
            row[f"ChipBit{i}"] = bit_label(cr["chip_bit"])
            if cr.get("judged"):
                row[f"Match_CH{i}"] = "✓" if cr["passed"] else "✗"
        elif i - 1 < len(ar.actual_currents):
            i_col = ar.actual_currents[i - 1]
            row[f"Chip_ICol{i}"] = f"{i_col:.6f}"
            chip_bit = chip_value_to_bit(i_col)
            row[f"ChipBit{i}"] = bit_label(chip_bit)
        if i - 1 < len(ar.actual_raws):
            row[f"Chip_ADC{i}"] = ar.actual_raws[i - 1]
    _fill_measured_infobits(row, ar.actual_currents, ar.chip_bits or None, mode=tc.mode)
    for i in range(1, 9):
        i_col = ar.actual_currents[i - 1] if i - 1 < len(ar.actual_currents) else 0.0
        row[f"CodedBit{i}"] = bit_label(chip_value_to_bit(i_col))
    return row


def _measured_row(tc: CaseRecord) -> dict:
    """实测行：RX hex、列电流、迭代；未测试时实测列留空。"""
    row: dict = {
        "RowKind": "实测行",
        "Attempt": "1" if has_measured(tc) else "",
        "CaseName": tc.name,
        "Mode": tc.mode,
        "SNR_dB": tc.snr_db,
        "FrameIdx": tc.frame_idx if tc.frame_idx is not None else "",
        "IterRef": tc.iter_ref if tc.iter_ref is not None else "",
    }

    for i in range(1, 13):
        row[f"In_CH{i}"] = ""
        row[f"In_Vrow{i}"] = ""

    sim = tc.sim_outputs or [0.0] * 8
    for i in range(1, 9):
        row[f"Sim_Out{i}"] = f"{sim[i - 1]:.6f}" if i - 1 < len(sim) else ""

    if not has_measured(tc):
        row.update(_blank_measured_fields())
        return row

    if tc.info_bits and tc.actual_currents and not tc.chip_bits:
        apply_bit_compare(tc)

    row.update(
        {
            "HardwareConverged": "Y" if tc.hardware_converged else "N",
            "FramePassed": "Y" if tc.passed else ("N" if tc.passed is False else ""),
            "ChannelMatch": f"{tc.match_count}/{compare_channels_for_mode(tc.mode)}",
            "Success": "Y" if tc.success else "N",
            "Error": tc.error,
            "RetryCount": tc.retries,
            "LowValueRetries": tc.low_value_retries,
            "Measured_Iterations": tc.iteration_count if tc.iteration_count else "",
            "SendHex": tc.tx_hex,
            "RecvHex": tc.rx_hex,
        }
    )

    ch_results = {r["channel"]: r for r in tc.channel_results}
    for i in range(1, 9):
        cr = ch_results.get(i)
        if cr:
            row[f"Chip_ICol{i}"] = f"{cr['chip_current']:.6f}"
            row[f"RealAdj{i}"] = f"{cr['chip_current']:.6f}"
            row[f"ChipBit{i}"] = bit_label(cr["chip_bit"])
            row[f"ChipSign{i}"] = bit_label(cr["chip_bit"])
            row[f"SimSign{i}"] = ""
            if cr.get("judged"):
                row[f"Match_CH{i}"] = "✓" if cr["passed"] else "✗"
            else:
                row[f"Match_CH{i}"] = "—"
        elif i - 1 < len(tc.actual_currents):
            i_col = tc.actual_currents[i - 1]
            row[f"Chip_ICol{i}"] = f"{i_col:.6f}"
            row[f"RealAdj{i}"] = f"{i_col:.6f}"
            chip_bit = chip_value_to_bit(i_col)
            row[f"ChipBit{i}"] = bit_label(chip_bit)
            row[f"ChipSign{i}"] = bit_label(chip_bit)
            row[f"SimSign{i}"] = ""
            row[f"Match_CH{i}"] = ""
        else:
            row[f"Chip_ICol{i}"] = ""
            row[f"RealAdj{i}"] = ""
            row[f"ChipSign{i}"] = ""
            row[f"SimSign{i}"] = ""
            row[f"Match_CH{i}"] = ""

        if i - 1 < len(tc.actual_raws):
            row[f"Chip_ADC{i}"] = tc.actual_raws[i - 1]
        else:
            row[f"Chip_ADC{i}"] = ""

    _fill_measured_infobits(row, tc.actual_currents, tc.chip_bits or None, mode=tc.mode)

    for i in range(1, 9):
        i_col = tc.actual_currents[i - 1] if i - 1 < len(tc.actual_currents) else 0.0
        row[f"CodedBit{i}"] = bit_label(chip_value_to_bit(i_col)) if has_measured(tc) else ""

    for i in range(5, 9):
        row.setdefault(f"InfoBit{i}", "")

    return row


def build_case_rows(cases: list[CaseRecord]) -> list[dict]:
    """每条用例：仿真行 + attempt_records 各次实测行；未测帧仅仿真行。"""
    rows: list[dict] = []
    for tc in cases:
        rows.append(_sim_row(tc))
        if tc.attempt_records:
            for ar in tc.attempt_records:
                rows.append(_measured_row_from_attempt(tc, ar))
        elif has_measured(tc):
            rows.append(_measured_row(tc))
    return rows


def build_snr_summary_rows(cases: list[CaseRecord]) -> list[dict]:
    points = analyze_by_snr(cases)
    rows: list[dict] = []
    for p in points:
        row: dict = {
            "SNR_dB": p.snr_db,
            "TotalFrames": p.case_count,
            "TestedFrames": p.tested_count,
            "ConvergedFrames": p.converged_count,
            "MeasuredBER": (
                f"{p.channel_ber:.6e}" if p.tested_count > 0 and math.isfinite(p.channel_ber) else ""
            ),
            "ErrorBits": p.error_bits,
            "TotalBits": p.channel_bits,
            "AvgIterations": f"{p.avg_iterations:.2f}",
            "StdIterations": f"{p.std_iterations:.2f}",
        }
        if p.avg_ref_iterations is not None:
            row["AvgRefIterations"] = f"{p.avg_ref_iterations:.2f}"
            if p.std_ref_iterations is not None:
                row["StdRefIterations"] = f"{p.std_ref_iterations:.2f}"
        if p.reference_ber is not None:
            row["ReferenceBER"] = f"{p.reference_ber:.6e}"
            row["ReferenceErrorBits"] = p.reference_error_bits
            row["ReferenceTotalBits"] = p.reference_channel_bits
        rows.append(row)
    return rows


def write_csv_report(path: str | Path, cases: list[CaseRecord]) -> None:
    rows = build_case_rows(cases)
    if not rows:
        raise ValueError("无测试用例可导出")
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()))
        writer.writeheader()
        writer.writerows(rows)


def write_excel_report(path: str | Path, cases: list[CaseRecord]) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    case_rows = build_case_rows(cases)
    if not case_rows:
        raise ValueError("无测试用例可导出")

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    fill_orange = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    fill_yellow = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    fill_sim = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    fill_meas = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")
    fill_retry = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
    font_ok = Font(color="006400")
    font_fail = Font(color="CC0000")

    wb = openpyxl.Workbook()
    ws_cases = wb.active
    ws_cases.title = "Cases"
    headers = list(case_rows[0].keys())
    ws_cases.append(headers)
    header_idx = {h: i + 1 for i, h in enumerate(headers)}

    case_by_name = {tc.name: tc for tc in cases}
    row_num = 2
    for row in case_rows:
        ws_cases.append([row[h] for h in headers])
        if row.get("RowKind") == "仿真行":
            for col in range(1, len(headers) + 1):
                ws_cases.cell(row=row_num, column=col).fill = fill_sim
        elif row.get("RowKind") == "实测行":
            attempt_str = str(row.get("Attempt", "")).strip()
            attempt_n = int(attempt_str) if attempt_str.isdigit() else 0
            row_fill = fill_retry if attempt_n > 1 else fill_meas
            for col in range(1, len(headers) + 1):
                ws_cases.cell(row=row_num, column=col).fill = row_fill
            tc = case_by_name.get(row.get("CaseName", ""))
            if tc and not tc.hardware_converged and frame_was_tested(tc):
                for col in range(1, len(headers) + 1):
                    ws_cases.cell(row=row_num, column=col).fill = fill_orange
            elif tc and tc.channel_results:
                for i in range(1, 5):
                    sim_key = f"Sim_Out{i}"
                    chip_key = f"Chip_ICol{i}"
                    if sim_key in header_idx and chip_key in header_idx:
                        cr = next((c for c in tc.channel_results if c["channel"] == i), None)
                        if cr and cr.get("judged") and not cr.get("passed"):
                            ws_cases.cell(row=row_num - 1, column=header_idx[sim_key]).fill = fill_yellow
                            ws_cases.cell(row=row_num, column=header_idx[chip_key]).fill = fill_yellow
            for i in range(1, 5):
                match_key = f"Match_CH{i}"
                if match_key in header_idx:
                    cell = ws_cases.cell(row=row_num, column=header_idx[match_key])
                    if cell.value == "✓":
                        cell.font = font_ok
                    elif cell.value == "✗":
                        cell.font = font_fail
        row_num += 1

    summary_rows = build_snr_summary_rows(cases)
    if summary_rows:
        ws_snr = wb.create_sheet("SNR_Summary")
        snr_headers = list(summary_rows[0].keys())
        ws_snr.append(snr_headers)
        for row in summary_rows:
            ws_snr.append([row[h] for h in snr_headers])

    wb.save(path)


def write_styled_dataset_xlsx(
    path: str | Path,
    fieldnames: list[str],
    rows: list[dict[str, str]],
    *,
    matrix_mode: int = 2,
    highlight_ber: bool = True,
    dataset_style: str | None = None,
) -> None:
    """导出测后交错 CSV 镜像为带背景色的 Excel。

    dataset_style:
      - full（①完整）：通信失败行橙色；通信成功但误码行浅黄；收敛无误码默认色
      - ok（②）：仅行底色
      - ber（③）：误码通道 CodeBit+Out 粉红高亮（行集应已 filter_ber_error_rows）
    """
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    style = dataset_style or ("ber" if highlight_ber else "ok")
    fill_sim = PatternFill(start_color="E3F2FD", end_color="E3F2FD", fill_type="solid")
    fill_meas = PatternFill(start_color="FFFDE7", end_color="FFFDE7", fill_type="solid")
    fill_retry = PatternFill(start_color="E8EAF6", end_color="E8EAF6", fill_type="solid")
    fill_comm_fail = PatternFill(start_color="FFF3E0", end_color="FFF3E0", fill_type="solid")
    fill_ber_row = PatternFill(start_color="FFF9C4", end_color="FFF9C4", fill_type="solid")
    fill_err = PatternFill(start_color="FFCDD2", end_color="FFCDD2", fill_type="solid")
    font_header = Font(bold=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Dataset"
    col_idx = {h: i + 1 for i, h in enumerate(fieldnames)}
    ws.append(fieldnames)
    for cell in ws[1]:
        cell.font = font_header

    mode = int(matrix_mode)
    coded_prefix = detect_coded_bit_prefix(fieldnames)
    pending_sim: dict[str, str] | None = None
    pending_sim_row: int | None = None

    def _highlight_cell(row_num: int, key: str) -> None:
        if key not in col_idx:
            return
        ws.cell(row=row_num, column=col_idx[key]).fill = fill_err

    for row_num, row in enumerate(rows, start=2):
        ws.append([row.get(h, "") for h in fieldnames])
        kind = str(row.get("RowKind", "")).strip()
        attempt_str = str(row.get("Attempt", "")).strip()
        attempt_n = int(attempt_str) if attempt_str.isdigit() else 0

        if kind == ROW_KIND_SIM or kind == "仿真行":
            pending_sim = row
            pending_sim_row = row_num
            for col in range(1, len(fieldnames) + 1):
                ws.cell(row=row_num, column=col).fill = fill_sim
            continue

        if kind == ROW_KIND_MEASURED or kind == "实测行":
            from csv_loader import measured_row_is_nonconverged

            row_fill = fill_retry if attempt_n > 1 else fill_meas
            if style == "full" and pending_sim is not None:
                if measured_row_is_nonconverged(row):
                    row_fill = fill_comm_fail
                elif measured_row_ber_mismatch_indices(
                    pending_sim, row, fieldnames, matrix_mode=mode
                ):
                    row_fill = fill_ber_row
            for col in range(1, len(fieldnames) + 1):
                ws.cell(row=row_num, column=col).fill = row_fill
            if attempt_n > 1 and "Attempt" in col_idx and row_fill in (
                fill_retry,
                fill_meas,
            ):
                ws.cell(row=row_num, column=col_idx["Attempt"]).fill = fill_retry

            if style == "ber" and pending_sim is not None and coded_prefix:
                for out_i, cb_i in measured_row_ber_mismatch_pairs(
                    pending_sim, row, fieldnames, matrix_mode=mode
                ):
                    _highlight_cell(row_num, f"Out[{out_i}]")
                    _highlight_cell(row_num, f"{coded_prefix}[{out_i}]")
                    if pending_sim_row is not None:
                        _highlight_cell(pending_sim_row, f"Out[{out_i}]")
                        _highlight_cell(pending_sim_row, f"{coded_prefix}[{cb_i}]")
            continue

        pending_sim = None
        pending_sim_row = None

    wb.save(path)
