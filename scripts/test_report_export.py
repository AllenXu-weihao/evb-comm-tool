"""测试报告导出冒烟测试（无需 GUI / 串口）。"""



from __future__ import annotations



import tempfile

from pathlib import Path

import sys



ROOT = Path(__file__).resolve().parents[1]

sys.path.insert(0, str(ROOT))



from analysis.aggregate import CaseRecord, apply_sim_compare, case_from_csv_group

from analysis.report_export import (

    build_case_rows,

    build_snr_summary_rows,

    write_csv_report,

    write_excel_report,

)

from csv_loader import load_test_csv





def _sample_case(name: str, snr: float, passed: bool) -> CaseRecord:

    tc = CaseRecord(

        name=name,

        mode=1,

        channels=[2500] * 12,

        sim_outputs=[0.15, -0.12, 0.08, -0.05, 0.0, 0.0, 0.0, 0.0],
        info_bits=[0, 1, 0, 1],

        snr_db=snr,

        iteration_count=12,

        actual_currents=[0.15, -0.12, 0.08, -0.05, 0.0, 0.0, 0.0, 0.0],

        actual_raws=[2650, 2380, 2550, 2480, 2500, 2500, 2500, 2500],

        tx_hex="AA " + "00 " * 12,

        rx_hex="0C " + "00 " * 16,

        hardware_converged=True,

        success=True,

    )

    apply_sim_compare(tc)

    if not passed:

        tc.actual_currents[0] = -0.15

        apply_sim_compare(tc)

    return tc





def main() -> None:

    cases = [_sample_case("1", 0.0, True), _sample_case("2", 0.0, False)]



    testdata = Path("data/测试集/mode1.csv")
    if not testdata.exists():
        testdata = Path("data/mode1.csv")
    if not testdata.exists():
        testdata = Path("data/测试集/testdata.csv")
    if not testdata.exists():
        testdata = Path("data/testdata.csv")

    if testdata.exists():

        groups = load_test_csv(testdata)

        cases = [case_from_csv_group(g) for g in groups.cases[:3]]

        for i, tc in enumerate(cases):

            tc.iteration_count = 10 + i

            info = tc.info_bits or [0, 0, 0, 0]
            currents = [0.0] * 8
            for ch in range(4):
                currents[ch] = -0.1 if info[ch] else 0.1
            tc.actual_currents = currents

            tc.actual_raws = [2600, 2400, 2550, 2480] + [2500] * 4

            tc.tx_hex = "AA 01"

            tc.rx_hex = "0A FF"

            tc.hardware_converged = True

            tc.success = True

            apply_sim_compare(tc)



    rows = build_case_rows(cases)

    assert rows, "expected at least one row"
    assert len(rows) == len(cases) * 2, "each case should export sim + measured rows"
    assert rows[0]["RowKind"] == "仿真行"
    assert rows[1]["RowKind"] == "实测行"
    assert "Sim_Out1" in rows[0]
    assert rows[0]["RecvHex"] == ""
    assert rows[1]["RecvHex"] != "" or rows[1]["RecvHex"] == ""
    assert "RealAdj1" in rows[1]
    assert "Chip_ICol1" in rows[1]
    assert "Match_CH1" in rows[1]
    assert "FramePassed" in rows[1]

    synthetic = [_sample_case("1", 0.0, True), _sample_case("2", 0.0, False)]
    syn_rows = build_case_rows(synthetic)
    pass_case = next(tc for tc in synthetic if tc.passed)
    assert pass_case.chip_bits == [0, 1, 0, 1]
    assert syn_rows[0]["InfoBit1"] == "0"
    assert syn_rows[0]["InfoBit2"] == "1"
    assert syn_rows[0]["InfoBit3"] == "0"
    assert syn_rows[0]["InfoBit4"] == "1"
    for i in range(1, 5):
        assert syn_rows[1][f"InfoBit{i}"] == "", f"measured InfoBit{i} should be blank"
        assert syn_rows[1][f"ChipBit{i}"] != "", f"measured ChipBit{i} should be set"



    summary = build_snr_summary_rows(cases)

    assert summary, "expected SNR summary when cases have measurements"



    with tempfile.TemporaryDirectory() as tmp:

        xlsx = Path(tmp) / "test_report.xlsx"

        csv_out = Path(tmp) / "test_report.csv"

        write_excel_report(xlsx, cases)

        write_csv_report(csv_out, cases)

        assert xlsx.stat().st_size > 0

        assert csv_out.stat().st_size > 0



        import openpyxl



        wb = openpyxl.load_workbook(xlsx)

        assert "Cases" in wb.sheetnames

        assert "SNR_Summary" in wb.sheetnames

        ws = wb["Cases"]

        headers = [c.value for c in ws[1]]

        assert "RowKind" in headers
        assert "Sim_Out1" in headers
        assert "RealAdj1" in headers
        assert "Chip_ICol1" in headers
        assert "Match_CH1" in headers
        assert "FramePassed" in headers



    print("OK: export smoke test passed")

    print(f"  case columns: {len(rows[0])}")

    print(f"  SNR summary rows: {len(summary)}")





if __name__ == "__main__":

    main()

