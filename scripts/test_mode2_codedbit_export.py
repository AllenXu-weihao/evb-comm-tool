"""Mode2 新数据集 mode2_0706：CodedBit 列、Out 硬判导出、CodeBit BER。

运行：py -3 scripts/test_mode2_codedbit_export.py
"""
from __future__ import annotations

import sys
import tempfile
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from analysis.aggregate import (
    analyze_by_snr,
    case_from_csv_group,
    count_measured_error_bits,
    count_sim_reference_error_bits,
)
from analysis.report_export import write_styled_dataset_xlsx
from analysis.sign_compare import (
    MODE2_BER_OUT_INDICES,
    chip_value_to_bit,
    mode2_coded_bits_from_outputs,
    sim_value_to_bit,
)
from csv_loader import (
    CODED_BIT_COUNT,
    detect_coded_bit_prefix,
    load_test_csv,
    resolve_export_fieldnames,
)


def _fixture_path() -> Path | None:
    for p in (
        ROOT / "data" / "mode2_0706.csv",
        Path(r"c:\Users\15209\Desktop\mode2_0706.csv"),
    ):
        if p.is_file():
            return p
    return None


def test_load_shape() -> None:
    path = _fixture_path()
    if path is None:
        print("SKIP: mode2_0706.csv not found")
        return

    loaded = load_test_csv(path)
    assert loaded.matrix_mode == 2
    assert len(loaded.cases) == 33000, f"expected 33000 cases, got {len(loaded.cases)}"
    assert loaded.dataset is not None

    prefix = detect_coded_bit_prefix(loaded.dataset.fieldnames)
    assert prefix == "CodedBit", loaded.dataset.fieldnames

    snr_counts = Counter(tc.snr_db for tc in [case_from_csv_group(g) for g in loaded.cases])
    assert len(snr_counts) == 11, f"expected 11 SNR points, got {sorted(snr_counts.keys())}"
    for snr, cnt in sorted(snr_counts.items(), key=lambda x: x[0]):
        assert cnt == 3000, f"SNR {snr}: expected 3000 frames, got {cnt}"

    points = analyze_by_snr([case_from_csv_group(g) for g in loaded.cases[:6000]])
    assert len(points) == 2
    assert points[0].reference_ber is not None
    assert points[0].reference_channel_bits > 0
    print(f"OK: load {len(loaded.cases)} cases, {len(snr_counts)} SNR × 3000 frames")


def test_codedbit_vs_out_reference_ber() -> None:
    """指定 CodeBit 可与 Sim Out 硬判不一致；参考 BER 统计二者差异。"""
    path = _fixture_path()
    if path is None:
        print("SKIP: mode2_0706.csv not found")
        return

    loaded = load_test_csv(path)
    assert loaded.dataset is not None
    prefix = detect_coded_bit_prefix(loaded.dataset.fieldnames) or "CodedBit"

    ber_mism = 0
    ber_total = 0
    tc_sample = [case_from_csv_group(g) for g in loaded.cases[:400]]
    for tc in tc_sample:
        sim = tc.sim_outputs or []
        coded = tc.coded_bits
        if not coded or len(sim) < max(MODE2_BER_OUT_INDICES) + 1:
            continue
        for i in MODE2_BER_OUT_INDICES:
            ber_total += 1
            expected = sim_value_to_bit(sim[i])
            actual = int(coded[i])
            if actual != expected:
                ber_mism += 1

    ref_err, ref_bits, n_ref = count_sim_reference_error_bits(tc_sample)
    assert n_ref == len(tc_sample)
    assert ref_bits == ber_total
    assert ref_err == ber_mism
    print(
        f"OK: reference BER on BER channels — "
        f"{ref_err}/{ref_bits} mismatches ({n_ref} sim rows)"
    )


def test_ber_channels_only_3567() -> None:
    path = _fixture_path()
    if path is None:
        print("SKIP: mode2_0706.csv not found")
        return

    loaded = load_test_csv(path)
    tc = case_from_csv_group(loaded.cases[0])
    sim = list(tc.sim_outputs or [])
    tc.actual_currents = list(sim)
    tc.iteration_count = 5
    tc.rx_hex = "00" * 17
    tc.hardware_converged = True
    assert count_measured_error_bits(tc) == 0

    idx = MODE2_BER_OUT_INDICES[0]
    if tc.coded_bits and len(tc.coded_bits) > idx:
        tc.coded_bits = list(tc.coded_bits)
        tc.coded_bits[idx] = sim_value_to_bit(sim[idx])
    tc.actual_currents[idx] = -sim[idx] if sim[idx] >= 0 else abs(sim[idx]) + 0.1
    assert count_measured_error_bits(tc) == 1

    ref_err, ref_bits, n_ref = count_sim_reference_error_bits([tc])
    assert n_ref == 1
    assert ref_bits == 4
    print(f"OK: measured BER flip 1 bit; reference BER err={ref_err}/{ref_bits}")


def test_export_preserves_codedbit() -> None:
    path = _fixture_path()
    if path is None:
        print("SKIP: mode2_0706.csv not found")
        return

    loaded = load_test_csv(path)
    assert loaded.dataset is not None
    ds = loaded.dataset
    cases = [case_from_csv_group(g) for g in loaded.cases[:4]]
    ds.rebuild_rows_from_cases(cases)

    prefix = detect_coded_bit_prefix(ds.fieldnames) or "CodedBit"
    sim_rows = [r for r in ds.rows if r.get("RowKind") == "仿真行"]
    assert len(sim_rows) >= 4
    for sim, tc in zip(sim_rows[:4], cases):
        for i in range(CODED_BIT_COUNT):
            expected = str(int(tc.coded_bits[i])) if tc.coded_bits else ""
            assert sim[f"{prefix}[{i}]"] == expected
        ber_bits = mode2_coded_bits_from_outputs(
            [float(sim[f"Out[{i}]"]) for i in range(8)]
        )
        assert len(ber_bits) == 4

    fn = resolve_export_fieldnames(ds.fieldnames)
    assert f"{prefix}[0]" in fn

    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "out.xlsx"
        write_styled_dataset_xlsx(xlsx, fn, ds.rows[:8], matrix_mode=2)
        import openpyxl

        wb = openpyxl.load_workbook(xlsx)
        headers = [c.value for c in wb.active[1]]
        assert f"{prefix}[0]" in headers

    print("OK: export keeps CodedBit columns and recomputes from Out")


def test_mode2_xlsx_highlights_codedbit_not_infobit() -> None:
    """Mode2 样式 Excel：误码红底仅 CodedBit[3,5,6,7]，InfoBit 不参与。"""
    from csv_loader import resolve_export_fieldnames

    fn = resolve_export_fieldnames(["CodedBit[0]"])
    prefix = "CodedBit"
    sim_row: dict[str, str] = {
        "RowKind": "仿真行",
        "Attempt": "",
        "InfoBit[0]": "0",
        "InfoBit[1]": "1",
        "InfoBit[2]": "0",
        "InfoBit[3]": "1",
        f"{prefix}[3]": "0",
        f"{prefix}[5]": "1",
    }
    meas_row: dict[str, str] = {
        "RowKind": "实测行",
        "Attempt": "1",
        "InfoBit[0]": "1",  # mismatch — must NOT highlight in Mode2
        "InfoBit[1]": "1",
        "InfoBit[2]": "0",
        "InfoBit[3]": "1",
        f"{prefix}[3]": "1",  # mismatch — must highlight
        f"{prefix}[5]": "1",
    }
    rows = [sim_row, meas_row]

    fill_err = "FFCDD2"

    with tempfile.TemporaryDirectory() as tmp:
        xlsx = Path(tmp) / "highlight.xlsx"
        write_styled_dataset_xlsx(xlsx, fn, rows, matrix_mode=2)
        import openpyxl

        wb = openpyxl.load_workbook(xlsx)
        ws = wb.active
        headers = {c.value: i + 1 for i, c in enumerate(ws[1])}

        def _fill_hex(row_idx: int, col_name: str) -> str:
            cell = ws.cell(row=row_idx, column=headers[col_name])
            rgb = getattr(cell.fill.start_color, "rgb", None) or ""
            return str(rgb)[-6:].upper()

        assert _fill_hex(2, "InfoBit[0]") != fill_err, "Mode2 must not highlight InfoBit"
        assert _fill_hex(3, "InfoBit[0]") != fill_err
        assert _fill_hex(2, f"{prefix}[3]") == fill_err
        assert _fill_hex(3, f"{prefix}[3]") == fill_err

    print("OK: Mode2 xlsx highlights CodedBit[3] only, not InfoBit")


def main() -> None:
    test_load_shape()
    test_codedbit_vs_out_reference_ber()
    test_ber_channels_only_3567()
    test_export_preserves_codedbit()
    test_mode2_xlsx_highlights_codedbit_not_infobit()


if __name__ == "__main__":
    main()
